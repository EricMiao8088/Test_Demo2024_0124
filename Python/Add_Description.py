import sys
ApiPath = r'C:\Program Files\ECU-TEST 2022.2\Templates\ApiClient'#ET安装路径下面的，找对应版本
sys.path.append(ApiPath)
from ApiClient import ApiClient
import openpyxl#需要安装openpyxl库
api = ApiClient()
import os

col_Description = 1#定义description的列数
col_test_name = 2#定义package的名字列数
description_list = []
pkg_name = []
def Add_Description(excelpath,packagepath):
    
    wb = openpyxl.load_workbook(excelpath)
    ws = wb[wb.sheetnames[0]]
    #print(ws)
    for row in range(ws.max_row-1):
        description = ws.cell(row = row + 2, column = col_Description).value#获取excel里面description内容
        test_name = ws.cell(row = row+2,column = col_test_name).value#获取package的name
        description_list.append(description)
        pkg_name.append(test_name)
        #print(description,test_name)
    for parent,direnames,filenames in os.walk(packagepath):#获取package所在根目录下面所有的文件夹，文件名
        #print(parent,'-----',direnames,'****',filenames,'++++')
        for name in filenames:
            if '.pkg' in name:#判断如果后缀是.pkg
                for li in pkg_name:#在pkg_name这个列表里面进行遍历，如果列表里面的名字一致，就添加description
                    if li + '.pkg' == name:
                        #print(li)
                        openpkg = api.PackageApi.OpenPackage(os.path.join(parent,name))#打开pkg
                        #print(os.path.join(parent,name))
                        openpkg.SetDescription(description_list[pkg_name.index(li)])#设置description
                        #print(description_list[pkg_name.index(li)],pkg_name.index(li))
                        openpkg.Save()

        



if __name__=="__main__":
    excelpath = r'C:\ET2022_4\1_Exercises\2_ECU-TEST_Advanced\Python\Description.xlsx'
    packagepath = r'C:\ET2022_2\WS03-ECU_TEST_Model_Testing\sldemo_autotrans\Packages\test'
    Add_Description(excelpath,packagepath)

