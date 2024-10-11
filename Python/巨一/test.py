import sys
import os
import openpyxl
ET_install_path = r'C:\Program Files\ECU-TEST 2022.4'#ET安装路径,需要修改
internalApi = os.path.join(ET_install_path,r'Templates\ExampleWorkspaces\Tutorial-ProcessSeparation\PythonPath')
sys.path.append(internalApi)
from ApiProxy import ApiProxy

api = ApiProxy()
mapping_file = api.ObjectApi.GlobalMappingApi.CreateMapping()

def mapping(path_excelpath):
    wb = openpyxl.load_workbook(path_excelpath)
    ws = wb[wb.sheetnames[0]]
    unit_path_list = []
    col_unit = 1
    
    model_mapping = api.DataBrowser.BrowseModel('Plant model').FindElements('*')
    for i in range(ws.max_row-1):
        unit_path = ws.cell(row = i + 2, column = col_unit).value#获取每个单元的路径，这个路径和ET里面的要相同
        for model_mapping_path in model_mapping:
            if unit_path in model_mapping_path:#先筛选出单个模块的路径，然后与对应的excel表里的信号名字拼凑起来
                unit_path_list.append(model_mapping_path)#将筛选出来的路径保存在一个列表里面
                #print(model_mapping_path)
        #print(unit_path_list)
        target_excel_name = unit_path.split('/')[-1]+'.xlsx'#打开对应单元的excel，获取信号名字
        wb_name = openpyxl.load_workbook(target_excel_name)
        ws_name = wb_name[wb_name.sheetnames[0]]
        
        for col in ws_name.iter_cols(max_row = 1):
            for cell in col:
                if cell.value != 'Input' and cell.value != 'Output':#获取对应excel表里的名字
                    #print(cell.value,'--------------')
                    for na in unit_path_list:
                        #print(cell.value,'-------')
                        if cell.value + '/<' in na:
                            referencename = cell.value
                            mappingpath = na
                            Type = 'Plant model'
                            #print(na)
                            category = unit_path.split('/')[-1]
                            Genmapping = mappingItem(Type,mappingpath,referencename,category)
    mapping_file.Save('model_mapping.xam')
def mappingItem(Type,mappingpath,referencename,Category):#这个函数，是用来定义ECU-TEST里面生成global mapping的步骤
    mappingItem = api.ObjectApi.PackageApi.MappingApi.CreateMappingItem(Type,mappingpath,referencename)
    mappingItem.SetCategory(Category)
    mapping_file.AddItem(mappingItem)

mapping(r'C:\ET2022_4\1_Exercises\2_ECU-TEST_Advanced\Python\巨一\test.xlsx')

