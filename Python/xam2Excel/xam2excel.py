import os
import os.path
import sys
import openpyxl#需要安装openpyxl的库
ApiPath = r'C:\Program Files\ECU-TEST 2022.4\Templates\ApiClient'#需要填写一下ApiClient的路径，如果安装在C盘下，只要改个ET的版本号
sys.path.append(ApiPath)
from ApiClient import ApiClient
api = ApiClient()
'''
此脚本需要在打开ECU-TEST的情况下运行。
'''
referencename_list = []
col_referenceName = 1
col_MappingType = 2
col_Sysidentify = 3
col_MappingTarget = 4
def xam2excel(excelpath,xamfile):#定义两个路径，一个是excel的路径，一个是xam的路径
    wb = openpyxl.load_workbook(excelpath)
    ws = wb[wb.sheetnames[0]]
    openxam = api.GlobalMappingApi.OpenMapping(xamfile)
    L1 = openxam.GetItems()
    for i in range(len(L1)):
        referencename = L1[i].GetReferenceName()
        targetpath = L1[i].GetTargetPath()
        modeltype = L1[i].GetAccessType()
        systemidentify = L1[i].GetSystemIdentifier()
        '''
        如果需要将Plant model从target中分离出来，就把下面注释的几行代码释放出来，然后将
        ws.cell(row = i + 2,column = col_MappingTarget).value = targetpath中的targetpath换成a即可
        '''
        # t1 = targetpath.split('/')[1:]
        # a = ("/".join(t1))
        ws.cell(row = i + 2, column = col_referenceName).value = referencename
        ws.cell(row = i + 2, column = col_MappingType).value = modeltype
        ws.cell(row = i + 2, column = col_Sysidentify).value = systemidentify
        ws.cell(row = i + 2,column = col_MappingTarget).value = targetpath
    wb.save(excelpath)
    


if __name__ == "__main__":
    xamfile = r'C:\ET2022_4\1_Exercises\2_ECU-TEST_Advanced\Parameters\model.xam'
    excelpath = r'C:\ET2022_4\1_Exercises\2_ECU-TEST_Advanced\Python\xam2excel.xlsx'
    xam2excel(excelpath,xamfile)
